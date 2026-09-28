# 以 v4.5 xlsx 的公式建立依賴圖，計算各「目標輸出」的前導格集合（精簡版至少要保留的格）
import openpyxl, re, sys, json
from openpyxl.formula import Tokenizer
from openpyxl.utils import range_boundaries, get_column_letter
from openpyxl.worksheet.formula import DataTableFormula
P='dist/20260926_CoreWeave收支模型_v4_5.xlsx'
wb=openpyxl.load_workbook(P)
def cells_of(sheet, ref):
    ref=ref.replace('$','')
    if ':' in ref:
        c1,r1,c2,r2=range_boundaries(ref)
    else:
        c1,r1,c2,r2=range_boundaries(ref+':'+ref)
    return [(sheet,r,c) for r in range(r1,r2+1) for c in range(c1,c2+1)]
deps={}; kind={}
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for x in row:
            v=x.value; k=(ws.title,x.row,x.column)
            if v is None: continue
            if isinstance(v,DataTableFormula):
                kind[k]='f'; deps[k]=None; continue
            if isinstance(v,str) and v.startswith('='):
                kind[k]='f'; ds=[]
                for t in Tokenizer(v).items:
                    if t.type=='OPERAND' and t.subtype=='RANGE':
                        s=t.value
                        if '!' in s:
                            sh,rf=s.rsplit('!',1); sh=sh.strip("'")
                        else: sh,rf=ws.title,s
                        try: ds+=cells_of(sh,rf)
                        except Exception: pass
                deps[k]=ds
            else:
                kind[k]='t' if isinstance(v,str) else 'n'
# 模擬運算表（輸入與假設 D166:E168）：每格＝以情境 1/2/3 重算整個模型，依賴公式列 D165:E165 與情境輸入格 C5
for r in (166,167,168):
    for c in (4,5):
        deps[('輸入與假設',r,c)]=[('輸入與假設',165,4),('輸入與假設',165,5),('輸入與假設',5,3),('輸入與假設',r,3)]
        kind[('輸入與假設',r,c)]='f'
def closure(targets):
    seen=set(); st=list(targets)
    while st:
        k=st.pop()
        if k in seen: continue
        seen.add(k)
        for d in deps.get(k) or []: st.append(d)
    return seen
def rng(sh,ref): return cells_of(sh,ref)
# 第 2 輪修訂：目標 4 改為季度追蹤整張；目標 2 加「總債務 ÷ EBITDA」；另加錨定矩陣與核心檢查
def rows(sh,rs): 
    ws=wb[sh]; return [(sh,r,c) for r in rs for c in range(1,ws.max_column+1) if kind.get((sh,r,c))=='f']
q=wb['季度追蹤']
T={
 'G1 三情境目標價與評等': rng('評價_DCF與目標價','C48:D70')+rng('摘要','C5:D15'),
 'G1b 錨定×倍數矩陣（穩健性一行）': rows('評價_DCF與目標價',range(74,86)),
 'G2 資金缺口與股權稀釋': rng('各期收支','C39:G41')+rng('各期收支','C58:G70')+rows('資產負債_新債與新股',[31]),
 'G3 共識與指引差異及原因': rng('摘要','C17:E48')+rng('摘要','C51:C54')+rows('摘要',[69]),
 'G4 季度追蹤（整張：焦點季第一屏＋實際數比較＋原因＋核對）': rows('季度追蹤',range(1,q.max_row+1))+rng('摘要','C61:G67'),
 'G5 反向 DCF 摘要': rng('摘要','C55:E58'),
 'G8 核心檢查（現金橋、季度核對）': rows('各期收支',range(74,85))+rows('檢查_連動',[5]),
}
allk=set()
for n,t in T.items():
    s=closure(t); allk|=s
    print(n,'公式',sum(1 for k in s if kind.get(k)=='f'),'數值輸入',sum(1 for k in s if kind.get(k)=='n'))
F=sum(1 for v in kind.values() if v=='f'); N=sum(1 for v in kind.values() if v=='n')
print('全部 公式',F,'數值',N,'｜聯集 公式',sum(1 for k in allk if kind.get(k)=='f'),'數值',sum(1 for k in allk if kind.get(k)=='n'))
per={}
for k,v in kind.items():
    d=per.setdefault(k[0],[0,0,0,0])
    if v=='f': d[0]+=1; d[1]+= k in allk
    if v=='n': d[2]+=1; d[3]+= k in allk
for s,d in per.items(): print(s,'公式 %d→%d 數值 %d→%d'%tuple(d))
rs={(k[0],k[1]) for k in allk if k in kind}
print('保留列',len(rs),'標籤字數',sum(len(v) for (s,r) in rs for v in [wb[s].cell(r,1).value] if isinstance(v,str) and not v.startswith('=')))
ws=wb['輸入與假設']
print('輸入列（數值）不在聯集：',[ (r,str(ws.cell(r,1).value)[:18]) for r in range(1,ws.max_row+1) if any(kind.get(('輸入與假設',r,c))=='n' for c in range(3,9)) and not any(('輸入與假設',r,c) in allk for c in range(3,9))])
