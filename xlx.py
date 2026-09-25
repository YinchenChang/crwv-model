import sys, json, shutil, subprocess
from openpyxl import load_workbook
src, sel, out = sys.argv[1], int(sys.argv[2]), sys.argv[3]
ak = int(sys.argv[4]) if len(sys.argv) > 4 else None  # 選填：EV/EBITDA 錨定年度
p=f'/home/claude/xl17_{sel}_{ak or 1}.xlsx'; shutil.copy(src,p)
wb=load_workbook(p); ws=wb['輸入與假設']
for r in range(1,ws.max_row+1):
    v=ws.cell(row=r,column=1).value
    if v and str(v).startswith('情境選擇'): ws.cell(row=r,column=3,value=sel)
    if ak and v and str(v).startswith('EV/EBITDA 錨定年度'): ws.cell(row=r,column=3,value=ak)
wb.save(p)
subprocess.run(['python3','/mnt/skills/public/xlsx/scripts/recalc.py',p,'120'],capture_output=True)
wb=load_workbook(p,data_only=True); o={}
for sh in wb.sheetnames:
    ws=wb[sh]
    for r in range(1,ws.max_row+1):
        n=ws.cell(row=r,column=1).value
        if not n: continue
        if sh=='評價_可比公司': v=[ws.cell(row=r,column=2+i).value for i in range(8)]
        elif sh=='檢查_連動': v=[ws.cell(row=r,column=2).value]
        else: v=[ws.cell(row=r,column=3+i).value for i in range(5)]
        if any(isinstance(x,(int,float)) for x in v): o[f"{sh}|{n}"]=v
json.dump(o,open(out,'w'),ensure_ascii=False); print('ok',sel,len(o))
