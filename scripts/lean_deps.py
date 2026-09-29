# 精簡版（PR #12）：以完整版 xlsx 的公式建立依賴圖，計算各「目標輸出」的前導格（精簡版至少要保留的格）。
# 第 3 輪：目標改以「分頁＋A 欄列名稱」定義（不寫死列號，完整版增刪列時不失準）；無法解析的參照直接報錯，不略過。
# 用法：python3 scripts/lean_deps.py [xlsx，預設 dist/ 內的 xlsx]（印出各目標的公式數與保留列數）
#   作為模組：from lean_deps import analyse; analyse(path) → dict（allk、keep、rows、kind、wb、targets…）；scripts/lean.py 使用
import sys, glob
import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.formula import DataTableFormula

IN, VA, CF, SM, QT = '輸入與假設', '評價_DCF與目標價', '各期收支', '摘要', '季度追蹤'


def analyse(path):
    wb = openpyxl.load_workbook(path)
    deps, kind = {}, {}

    def cells_of(sheet, ref):
        ref = ref.replace('$', '')
        c1, r1, c2, r2 = range_boundaries(ref if ':' in ref else ref + ':' + ref)
        return [(sheet, r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]

    dt_cells = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for x in row:
                v = x.value; k = (ws.title, x.row, x.column)
                if v is None: continue
                if isinstance(v, DataTableFormula):
                    kind[k] = 'f'; dt_cells.append((k, v)); continue
                if isinstance(v, str) and v.startswith('='):
                    kind[k] = 'f'; ds = []
                    for t in Tokenizer(v).items:
                        if t.type == 'OPERAND' and t.subtype == 'RANGE':
                            s = t.value
                            sh, rf = (s.rsplit('!', 1) if '!' in s else (ws.title, s))
                            sh = sh.strip("'")
                            if sh not in wb.sheetnames: raise ValueError(f'{k}：參照不存在的分頁 {s}')
                            ds += cells_of(sh, rf)   # 解析失敗即報錯
                    deps[k] = ds
                else:
                    kind[k] = 't' if isinstance(v, str) else 'n'
    # 模擬運算表：區塊內每格＝以輸入值代入輸入格、重算公式列；依賴公式列、輸入格與該列輸入值
    for (k, v) in dt_cells:
        s = k[0]; c1, r1, c2, r2 = range_boundaries(v.ref); inp = cells_of(s, v.r1)[0]
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                kind[(s, r, c)] = 'f'
                deps[(s, r, c)] = [(s, r1 - 1, cc) for cc in range(c1, c2 + 1)] + [inp, (s, r, c1 - 1)]

    def rows_where(sheet, pred):
        ws = wb[sheet]
        return [r for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, str) and pred(ws.cell(r, 1).value)]

    def row_of(sheet, label, nth=0):
        rs = rows_where(sheet, lambda a: a == label)
        if len(rs) <= nth: raise KeyError(f'找不到列：{sheet}｜{label}')
        return rs[nth]

    def cells(sheet, rs):
        ws = wb[sheet]
        return [(sheet, r, c) for r in rs for c in range(2, ws.max_column + 1) if kind.get((sheet, r, c)) == 'f']

    def span(sheet, a, b):
        return list(range(row_of(sheet, a), row_of(sheet, b) + 1))

    T = {
        'G1 三情境目標價與評等': cells(VA, span(VA, 'DCF 每股（推到目標價時點）', '目標價區間｜判斷句')) + cells(SM, rows_where(SM, lambda a: a.startswith('結論｜'))),
        'G1b 錨定×倍數矩陣': cells(VA, rows_where(VA, lambda a: a.startswith('錨定×倍數｜'))),
        'G2 資金缺口與股權稀釋': cells(CF, span(CF, '營運缺口（不含融資、不含還本）', '期末累積現金') + span(CF, '新債舉借', '期末總債務（既有＋可轉債＋新債＋高息債）'))
                              + cells('資產負債_新債與新股', rows_where('資產負債_新債與新股', lambda a: a.startswith('總債務 ÷ EBITDA')))
                              + cells('檢查_連動', rows_where('檢查_連動', lambda a: 'FY28 起收入是否依賴未簽約' in a)),
        'G3 共識與指引差異及原因': cells(SM, rows_where(SM, lambda a: a.startswith(('差異', '隱含｜', '驗證｜目標價')))),
        'G4 季度追蹤（整張）': cells(QT, range(1, wb[QT].max_row + 1)) + cells(SM, rows_where(SM, lambda a: a.startswith('驗證｜') and '評等分布' not in a)),
        'G5 反向 DCF 摘要': cells(SM, rows_where(SM, lambda a: a.startswith('反向 DCF｜'))),
        'G1b／G11 穩健性與關鍵假設（建置時快照）': cells(VA, rows_where(VA, lambda a: a.startswith(('穩健性｜', '關鍵假設｜', '快照｜')))),
        'G8 核心檢查（現金橋）': cells(CF, span(CF, '四層橋（驗算）', '差異（應為 0）')) + cells('檢查_連動', rows_where('檢查_連動', lambda a: a == '現金橋差異')),
    }
    T = {k: [c for c in v if c in kind] for k, v in T.items()}

    def closure(targets):
        seen, st = set(), list(targets)
        while st:
            k = st.pop()
            if k in seen: continue
            seen.add(k); st += deps.get(k) or []
        return seen

    per_goal, allk = {}, set()
    for n, t in T.items():
        if not t: raise KeyError(f'{n}：目標沒有任何公式格（列名稱改了？）')
        s = closure(t); allk |= s
        per_goal[n] = (sum(1 for k in s if kind.get(k) == 'f'), sum(1 for k in s if kind.get(k) == 'n'))
    # 列層級的不動點：精簡版以整列為單位保留，保留列上的每個公式格（即使本身不在目標依賴鏈內）的前導格也要保留，
    # 否則刪列後該格會變成 #REF!（第 3 輪可行性測試發現：損益「全年 EBITDA」FY26 欄）。被引用的空白格（例如未公布的實際數）也保留。
    keep = set(allk)
    while True:
        rows = {(s, r) for (s, r, c) in keep if s in wb.sheetnames}
        extra = [k for k in kind if kind[k] == 'f' and (k[0], k[1]) in rows and k not in keep]
        if not extra: break
        keep |= closure(extra)
    rows = {(s, r) for (s, r, c) in keep if s in wb.sheetnames}
    return dict(wb=wb, kind=kind, deps=deps, allk=allk, keep=keep, rows=rows, targets=T, per_goal=per_goal, closure=closure)


if __name__ == '__main__':
    path = sys.argv[1] if len(sys.argv) > 1 else glob.glob('dist/*.xlsx')[0]
    g = analyse(path); kind, allk = g['kind'], g['allk']
    for n, (f, i) in g['per_goal'].items(): print(f'{n}：公式 {f}、數值輸入 {i}')
    F = sum(1 for v in kind.values() if v == 'f'); N = sum(1 for v in kind.values() if v == 'n')
    print(f'全部：公式 {F}、數值 {N}｜聯集：公式 {sum(1 for k in allk if kind.get(k) == "f")}、數值 {sum(1 for k in allk if kind.get(k) == "n")}')
    per = {}
    for k, v in kind.items():
        d = per.setdefault(k[0], [0, 0, 0, 0])
        if v == 'f': d[0] += 1; d[1] += k in allk
        if v == 'n': d[2] += 1; d[3] += k in allk
    for s, d in per.items(): print(f'  {s}：公式 {d[0]}→{d[1]}、數值 {d[2]}→{d[3]}')
    kp = g['keep']
    print(f'整列保留（不動點）後：公式 {sum(1 for k in kp if kind.get(k) == "f")}、數值 {sum(1 for k in kp if kind.get(k) == "n")}、列 {len(g["rows"])}')
